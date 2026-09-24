"""Exports for presentations, styled like RosterResource: Excel, CSV zip, printable HTML."""
from __future__ import annotations

import html
import io
import re
import zipfile

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import config, fgstyle

ATTRIBUTION = "Source: FanGraphs RosterResource (fangraphs.com)"
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(title: str, used: set[str]) -> str:
    base = (_BAD_SHEET_CHARS.sub("-", title).strip("'") or "Sheet")[:31]
    name, n = base, 2
    while name.lower() in {u.lower() for u in used}:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def _slug(title: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")[:60] or "table"


def _hex(c: str) -> str:
    return c.lstrip("#").upper()


def to_excel(items: list[tuple[str, pd.DataFrame]], meta: dict) -> bytes:
    t = fgstyle.TOKENS
    thin = Side(style="thin", color=_hex(t["border"]))
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name="Arial", bold=True, size=10)
    head_fill = PatternFill("solid", fgColor=_hex(t["header_bg"]))
    title_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor=_hex(t["green"]))
    body_font = Font(name="Arial", size=10)
    alt_fill = PatternFill("solid", fgColor=_hex(t["row_alt"]))
    acq_fill = PatternFill("solid", fgColor=_hex(t["acquired_bg"]))
    il_fill = PatternFill("solid", fgColor=_hex(t["il_bg"]))
    season = config.current_season()

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        used: set[str] = {"About"}
        about = pd.DataFrame(
            [("Report", meta.get("title", "")), ("Data snapshot", meta.get("snapshot", "")),
             ("Exported", meta.get("exported", "")), ("Source", ATTRIBUTION)]
            + [(f"Sheet: {t_}", f"{len(df)} rows") for t_, df in items],
            columns=["Field", "Value"])
        about.to_excel(xw, sheet_name="About", index=False)
        ws = xw.sheets["About"]
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 40, 60
        for c in ws[1]:
            c.font, c.fill, c.border = head_font, head_fill, border

        for title, df in items:
            name = _sheet_name(title, used)
            out = df.copy()
            out.columns = [fgstyle.header_label(c) for c in out.columns]
            # Row 1: green section bar with the panel title, row 2: headers, data from row 3.
            out.to_excel(xw, sheet_name=name, index=False, startrow=1)
            ws = xw.sheets[name]
            ncol = max(1, out.shape[1])
            ws.cell(1, 1, f"{title.upper()}  ({ATTRIBUTION})")
            for j in range(1, ncol + 1):
                ws.cell(1, j).fill = title_fill
                ws.cell(1, j).font = title_font
            ws.freeze_panes = "A3"
            for c in ws[2]:
                c.font, c.fill, c.border = head_font, head_fill, border
                c.alignment = Alignment(horizontal="center")
            by_norm = {fgstyle.norm(c): c for c in df.columns}
            for i, (_, row) in enumerate(df.iterrows()):
                cls = fgstyle.row_class(row, by_norm, season)
                fill = acq_fill if cls == "fg-acq" else il_fill if cls == "fg-il" else (alt_fill if i % 2 else None)
                for j, col in enumerate(df.columns, 1):
                    cell = ws.cell(i + 3, j)
                    cell.font, cell.border = body_font, border
                    if fill:
                        cell.fill = fill
                    if isinstance(row[col], (int, float)) and not isinstance(row[col], bool):
                        cell.alignment = Alignment(horizontal="center")
                        if fgstyle._MONEY_COLS.search(str(col)):
                            cell.number_format = '"$"#,##0'
            for j, col in enumerate(out.columns, 1):
                sample = [str(col)] + ["" if pd.isna(v) else str(v) for v in out.iloc[:, j - 1].head(200)]
                ws.column_dimensions[ws.cell(2, j).column_letter].width = min(max(map(len, sample)) + 3, 50)
    return buf.getvalue()


def to_csv_zip(items: list[tuple[str, pd.DataFrame]], meta: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        seen: set[str] = set()
        for i, (title, df) in enumerate(items, 1):
            name = f"{i:02d}_{_slug(title)}.csv"
            while name in seen:
                name = name.replace(".csv", "_x.csv")
            seen.add(name)
            zf.writestr(name, df.to_csv(index=False))
        zf.writestr("README.txt", f"{meta.get('title','')}\nSnapshot: {meta.get('snapshot','')}\n"
                                  f"Exported: {meta.get('exported','')}\n{ATTRIBUTION}\n")
    return buf.getvalue()


def to_html(items: list[tuple[str, pd.DataFrame]], meta: dict) -> str:
    esc = html.escape
    blocks = "\n".join(
        f"<div class='fg-block'>{fgstyle.section_html(title, f'{len(df)} rows')}"
        f"{fgstyle.table_html(df, max_rows=100000)}"
        f"<div class='fg-src'>{esc(ATTRIBUTION)}. Snapshot {esc(meta.get('snapshot', ''))}.</div></div>"
        for title, df in items)
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(meta.get('title', 'Roster report'))}</title>
<style>{fgstyle.report_css()}</style></head><body>
{fgstyle.page_title_html(meta.get('title', 'Roster report'),
                         f"Data snapshot {meta.get('snapshot', '')}. Exported {meta.get('exported', '')}.")}
{blocks}
</body></html>"""
